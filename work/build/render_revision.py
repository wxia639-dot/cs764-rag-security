from bind_revised_results import *
import re,shutil,html
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
OUT=ROOT/'outputs';T=ROOT/'work/ieee_report';F=T/'figures'
REPO='https://github.com/wxia639-dot/cs764-rag-security'
EQS={'probe':r'\begin{align}p(w)&=\sigma\!\left(\mathbf a^{\mathsf T}\mathbf h(w)+b\right),\\p(d)&=\max_{w\in W(d)}p(w).\end{align}',
'combined':r'\begin{equation}s_{\lambda}(d)=1-\bigl(1-p(d)\bigr)\bigl(1-\lambda r(d)\bigr).\end{equation}',
'threshold':r'\begin{equation}\operatorname{keep}_{m}(d)=\mathbf{1}\!\left[s_m(d)<\tau_m\right].\end{equation}'}
CAPS={'pipeline':'Executed request path and source-level attack insertion. The dashed line marks the boundary at which retrieved evidence is admitted to generation. Empty retained context returns UNKNOWN.',
'curves':'Held-out passage ROC and precision--recall curves from stored scores. The sample contains 111 attacked and 37 clean records. Curves describe ranking; fixed operating points appear in Table II.',
'sensitivity':'Cached-score sensitivity on the original passage test. Left: each weight is calibrated on audited validation. Right: descriptive fixed-weight threshold sweep. No test-selected setting is deployed.',
'tradeoffs':'Expanded security and QA outcomes (first three panels) alongside gate cost from the separate uncached 16-email pilot (fourth panel). The panels use different, explicitly labelled samples; expanded execution reuse is excluded from timing.'}
def esc(s):
    for a,b in [('\\',r'\textbackslash{}'),('&',r'\&'),('%',r'\%'),('$',r'\$'),('#',r'\#'),('_',r'\_'),('{',r'\{'),('}',r'\}')]:
        if a=='\\':continue
        s=s.replace(a,b)
    s=s.replace('–','--').replace('—','---').replace('→',r'$\rightarrow$')
    s=re.sub(r'\[(\d+)\]',lambda m:r'\cite{ref'+m[1]+'}',s)
    return s
def prose(t):
    if 'CODE_REPOSITORY_FOOTNOTE' in t:
        left,right=t.split('CODE_REPOSITORY_FOOTNOTE');return esc(left)+r'\footnote{Project code and raw results: \url{'+REPO+r'}. The source archive preserves the original directory structure.}'+esc(right)
    return esc(t)
def tabletex(key):
    cap,rows=TABLES[key];n=len(rows[0]);width='0.33' if n<=3 else '0.30'
    spec='p{'+width+r'\columnwidth}'+'r'*(n-1)
    lines=[r'\begin{table}[!t]',r'\caption{'+esc(cap)+'}',r'\label{tab:'+key+'}',r'\centering\footnotesize',r'\setlength{\tabcolsep}{3pt}',r'\begin{tabular}{'+spec+'}',r'\toprule']
    for i,row in enumerate(rows):
        # Long detector names wrap inside the first column.
        lines.append(' & '.join(esc(str(x)) for x in row)+r' \\')
        if i==0:lines.append(r'\midrule')
    return '\n'.join(lines+[r'\bottomrule\end{tabular}',r'\end{table}'])
tex=[r'\documentclass[conference,a4paper]{IEEEtran}',r'\usepackage{fontspec}',r'\setmainfont{Times New Roman}',r'\usepackage{amsmath,amssymb,graphicx,booktabs,url}',r'\usepackage[hidelinks]{hyperref}',r'\usepackage{xurl}',r'\urlstyle{same}',r'\begin{document}',r'\title{'+TITLE+'}',r'\author{\IEEEauthorblockN{'+AUTHORS+r'}\IEEEauthorblockA{University of Auckland\\COMPSCI 764}}',r'\maketitle',r'\begin{abstract}',esc(ABSTRACT),r'\end{abstract}',r'\begin{IEEEkeywords}retrieval-augmented generation, indirect prompt injection, passage filtering, calibration\end{IEEEkeywords}']
for title,items in SECTIONS:
    tex.append(r'\section{'+esc(title)+'}')
    for h,t in items:
        if h=='EQ':tex.append(EQS[t]);continue
        if h=='TABLE':tex.append(tabletex(t));continue
        if h=='FIG':
            wide=t=='tradeoffs';env='figure*' if wide else 'figure';w=r'\textwidth' if wide else r'\columnwidth'
            tex.extend([r'\begin{'+env+'}[!t]',r'\centering',r'\includegraphics[width='+w+']{figures/'+t+'.pdf}',r'\caption{'+esc(CAPS[t])+'}',r'\label{fig:'+t+'}',r'\end{'+env+'}']);continue
        if h:tex.append(r'\subsection{'+esc(h)+'}')
        tex.append(prose(t)+'\n')
tex.extend([r'\clearpage',r'\begin{thebibliography}{99}'])
for i,ref in enumerate(REFERENCES,1):
    rr=re.sub(r'^\[\d+\]\s*','',ref)
    # Use typographic double quotes around titles, italics for venues, and breakable URLs.
    match=re.search(r"'(.+)'",rr)
    if match:rr=rr[:match.start()]+'TITLESTART'+match[1]+'TITLEEND'+rr[match.end():]
    parts=re.split(r'(https?://\S+)',rr);formatted=''
    for part in parts:
        if part.startswith('http'):formatted+=r'\url{'+part.rstrip('.')+'}'+('.' if part.endswith('.') else '')
        else:formatted+=esc(part)
    formatted=formatted.replace('TITLESTART','``').replace('TITLEEND',"''")
    formatted=re.sub(r'(in )([^,]+)',lambda m:m[1]+r'\emph{'+m[2]+'}',formatted)
    tex.append(r'\bibitem{ref'+str(i)+'} '+formatted)
tex.extend([r'\end{thebibliography}',r'\end{document}'])
(T/'main.tex').write_text('\n'.join(tex),encoding='utf8')
(ROOT/'work/build/revised_manuscript.json').write_text(json.dumps({'title':TITLE,'abstract':ABSTRACT,'sections':SECTIONS,'references':REFERENCES,'tables':TABLES,'captions':CAPS},indent=2),encoding='utf8')

# Editable companion: same narrative, tables, native Word equations and figures.
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929);sec.top_margin=Inches(.75);sec.bottom_margin=Inches(1);sec.left_margin=Inches(.625);sec.right_margin=Inches(.625)
for sty in ['Normal','Title','Heading 1','Heading 2']:
    ss=doc.styles[sty];ss.font.name='Times New Roman';ss.font.color.rgb=RGBColor(0,0,0);ss.font.size=Pt(10 if sty=='Normal' else 11)
    ss.paragraph_format.space_after=Pt(6)
    pp=ss.element.find(qn('w:pPr'))
    if pp is not None:
        for border in list(pp.findall(qn('w:pBdr'))):pp.remove(border)
doc.styles['Title'].font.size=Pt(20)
doc.add_paragraph(TITLE,'Title').alignment=1;doc.add_paragraph(AUTHORS+'\nUniversity of Auckland | COMPSCI 764').alignment=1
doc.add_paragraph('Abstract','Heading 1');doc.add_paragraph(ABSTRACT)
sec=doc.add_section(WD_SECTION_START.CONTINUOUS);cols=sec._sectPr.find(qn('w:cols'));cols.set(qn('w:num'),'2');cols.set(qn('w:space'),'180')
plain_eq={'probe':'p(w) = σ(aᵀh(w) + b)     p(d) = max₍w∈W(d)₎ p(w)', 'combined':'sλ(d) = 1 − (1 − p(d))(1 − λr(d))', 'threshold':'keepₘ(d) = 𝟙[sₘ(d) < τₘ]'}
mathml={
'probe':'<mrow><mi>p</mi><mo>(</mo><mi>w</mi><mo>)</mo><mo>=</mo><mi>σ</mi><mo>(</mo><msup><mi mathvariant="bold">a</mi><mi mathvariant="normal">T</mi></msup><mi mathvariant="bold">h</mi><mo>(</mo><mi>w</mi><mo>)</mo><mo>+</mo><mi>b</mi><mo>)</mo><mo>,</mo><mspace width="1em"/><mi>p</mi><mo>(</mo><mi>d</mi><mo>)</mo><mo>=</mo><munder><mo>max</mo><mrow><mi>w</mi><mo>∈</mo><mi>W</mi><mo>(</mo><mi>d</mi><mo>)</mo></mrow></munder><mi>p</mi><mo>(</mo><mi>w</mi><mo>)</mo></mrow>',
'combined':'<mrow><msub><mi>s</mi><mi>λ</mi></msub><mo>(</mo><mi>d</mi><mo>)</mo><mo>=</mo><mn>1</mn><mo>−</mo><mo>(</mo><mn>1</mn><mo>−</mo><mi>p</mi><mo>(</mo><mi>d</mi><mo>)</mo><mo>)</mo><mo>(</mo><mn>1</mn><mo>−</mo><mi>λ</mi><mi>r</mi><mo>(</mo><mi>d</mi><mo>)</mo><mo>)</mo></mrow>',
'threshold':'<mrow><msub><mi mathvariant="normal">keep</mi><mi>m</mi></msub><mo>(</mo><mi>d</mi><mo>)</mo><mo>=</mo><mi>𝟙</mi><mo>[</mo><msub><mi>s</mi><mi>m</mi></msub><mo>(</mo><mi>d</mi><mo>)</mo><mo>&lt;</mo><msub><mi>τ</mi><mi>m</mi></msub><mo>]</mo></mrow>'}
math_transform=etree.XSLT(etree.parse('C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL'))
figno=0;tabno=0
for title,items in SECTIONS:
    doc.add_paragraph(title,'Heading 1')
    for h,t in items:
        if h=='EQ':
            p=doc.add_paragraph();m=math_transform(etree.fromstring(('<math xmlns="http://www.w3.org/1998/Math/MathML">'+mathml[t]+'</math>').encode('utf8'))).getroot();p._p.append(m);continue
        if h=='FIG':
            figno+=1;doc.add_picture(str(F/(t+'.png')),width=Inches(3.25));doc.add_paragraph(f'Figure {figno}. '+CAPS[t]);continue
        if h=='TABLE':
            tabno+=1;cap,rows=TABLES[t];doc.add_paragraph(f'Table {tabno}. '+cap);tb=doc.add_table(rows=0,cols=len(rows[0]));tb.style='Table Grid'
            for row in rows:
                cells=tb.add_row().cells
                for cell,txt in zip(cells,row):cell.text=str(txt)
            for row in tb.rows:
                for c in row.cells:
                    for p in c.paragraphs:
                        for run in p.runs:run.font.size=Pt(7)
            continue
        if h:doc.add_paragraph(h,'Heading 2')
        doc.add_paragraph(t.replace('CODE_REPOSITORY_FOOTNOTE',REPO))
doc.add_page_break();doc.add_paragraph('References','Heading 1')
for ref in REFERENCES:
    p=doc.add_paragraph();parts=re.split(r"('.*')",ref,maxsplit=1)
    for part in parts:
        if not part.startswith("'") and 'in ' in part:
            before,after=part.split('in ',1);venue,sep,tail=after.partition(',')
            for txt,italic in [(before+'in ',False),(venue,True),(sep+tail,False)]:
                rr=p.add_run(txt);rr.italic=italic;rr.font.size=Pt(8)
        else:
            rr=p.add_run('“'+part[1:-1]+'”' if part.startswith("'") else part);rr.font.size=Pt(8)
doc.save(OUT/'CS764_RAG_Security_Report_Revised.docx')
print(json.dumps({'abstract_words':len(ABSTRACT.split()),'body_words':sum(len(t.split()) for _,items in SECTIONS for h,t in items if h not in ['EQ','FIG','TABLE']),'references':len(REFERENCES),'sections':len(SECTIONS)}))
